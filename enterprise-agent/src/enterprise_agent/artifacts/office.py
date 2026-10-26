"""Portable Office renderers without a private Node runtime."""

import io
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from enterprise_agent.artifacts.models import ArtifactSpec
from enterprise_agent.artifacts.renderers import ArtifactRenderer, RenderedArtifact
from enterprise_agent.artifacts.templates import CompanyTemplate


class XlsxRenderer(ArtifactRenderer):
    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        book = Workbook()
        summary = book.active
        summary.title = "Summary"
        summary.append([spec.title])
        if spec.subtitle:
            summary.append([spec.subtitle])
        for section in spec.sections:
            summary.append([section.heading])
            for text in [section.body, *section.bullets]:
                for start in range(0, len(text) or 1, 600):
                    summary.append([text[start : start + 600]])
        for index, table in enumerate(spec.tables, 1):
            sheet = book.create_sheet(f"Table {index}")
            sheet.append(table.columns)
            for row in table.rows:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        if spec.sources:
            sheet = book.create_sheet("Sources")
            sheet.append(["Source", "Reference"])
            for source in spec.sources:
                sheet.append([source.label, source.reference])
        for sheet in book:
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # Never evaluate enterprise text as a formula.
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=template.primary_color.lstrip("#"))
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(
                    70, max(16, max(len(str(c.value or "")) for c in column) + 2)
                )
            for row in sheet:
                lines = max(
                    math.ceil(
                        len(str(cell.value or ""))
                        / sheet.column_dimensions[cell.column_letter].width
                    )
                    + str(cell.value or "").count("\n")
                    for cell in row
                )
                sheet.row_dimensions[row[0].row].height = min(409, max(20, lines * 15 + 8))
        output = io.BytesIO()
        book.save(output)
        return RenderedArtifact(
            output.getvalue(),
            "xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


class PptxRenderer(ArtifactRenderer):
    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

        def slide(title: str, lines: list[str]):
            page = deck.slides.add_slide(deck.slide_layouts[6])
            heading = page.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.8))
            heading.text_frame.text = title
            heading.text_frame.paragraphs[0].font.size = Pt(28)
            heading.text_frame.paragraphs[0].font.bold = True
            heading.text_frame.paragraphs[0].font.name = template.font_family
            heading.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(
                template.primary_color.lstrip("#")
            )
            body = page.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
            body.text_frame.word_wrap = True
            for index, line in enumerate(lines):
                p = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
                p.text, p.font.size, p.font.name = line, Pt(17), template.font_family
            return page

        def paginate(title: str, lines: list[str]) -> None:
            chunks: list[list[str]] = [[]]
            for line in lines:
                for start in range(0, len(line) or 1, 650):
                    piece = line[start : start + 650]
                    if sum(map(len, chunks[-1])) + len(piece) > 850 or len(chunks[-1]) >= 7:
                        chunks.append([])
                    chunks[-1].append(piece)
            for chunk in chunks:
                slide(title, chunk)

        slide(spec.title, [spec.subtitle or template.company_name])
        for section in spec.sections:
            paginate(section.heading, ([section.body] if section.body else []) + section.bullets)
        for table in spec.tables:
            chars_per_line = max(15, int(11.7 / len(table.columns) * 11))
            pages, heights, used = [[]], [[]], 0.0
            for row in table.rows:
                height = max(
                    0.45,
                    max(math.ceil(len(str(v or "")) / chars_per_line) for v in row) * 0.21 + 0.18,
                )
                if pages[-1] and (used + height > 4.4 or len(pages[-1]) >= 6):
                    pages.append([])
                    heights.append([])
                    used = 0.0
                pages[-1].append(row)
                heights[-1].append(height)
                used += height
            for rows, sizes in zip(pages, heights, strict=True):
                page = slide(table.title or "Data", [])
                grid = page.shapes.add_table(
                    len(rows) + 1,
                    len(table.columns),
                    Inches(0.8),
                    Inches(1.5),
                    Inches(11.7),
                    Inches(0.6 + sum(sizes)),
                ).table
                grid.rows[0].height = Inches(0.6)
                for index, height in enumerate(sizes, 1):
                    grid.rows[index].height = Inches(height)
                for r, values in enumerate([table.columns, *rows]):
                    for c, value in enumerate(values):
                        cell = grid.cell(r, c)
                        cell.text = "" if value is None else str(value)
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = Pt(12)
                            paragraph.font.name = template.font_family
                            paragraph.font.bold = r == 0
        if spec.sources:
            paginate("Sources", [f"{s.label}: {s.reference}" for s in spec.sources])
        output = io.BytesIO()
        deck.save(output)
        return RenderedArtifact(
            output.getvalue(),
            "pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
