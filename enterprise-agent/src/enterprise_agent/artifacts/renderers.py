"""Artifact renderers for text, Office, and PDF formats."""

import json
import logging
import subprocess
import tempfile
from abc import ABC, abstractmethod
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from enterprise_agent.artifacts.models import ArtifactFormat, ArtifactSpec
from enterprise_agent.artifacts.templates import CompanyTemplate

logger = logging.getLogger(__name__)


class ArtifactRenderingError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class RenderedArtifact:
    content: bytes
    extension: str
    media_type: str


class ArtifactRenderer(ABC):
    @abstractmethod
    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact: ...


class MarkdownRenderer(ArtifactRenderer):
    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        lines = [f"# {spec.title}"]
        if spec.subtitle:
            lines.extend(["", spec.subtitle])
        lines.extend(["", f"_Prepared for {template.company_name}_"])
        for section in spec.sections:
            lines.extend(["", f"## {section.heading}"])
            if section.body:
                lines.extend(["", section.body])
            lines.extend(f"- {bullet}" for bullet in section.bullets)
        for table in spec.tables:
            if table.title:
                lines.extend(["", f"## {table.title}"])
            lines.extend(
                [
                    "",
                    "| " + " | ".join(table.columns) + " |",
                    "| " + " | ".join("---" for _ in table.columns) + " |",
                ]
            )
            for row in table.rows:
                cells = [str(value if value is not None else "").replace("|", "\\|") for value in row]
                lines.append("| " + " | ".join(cells) + " |")
        if spec.sources:
            lines.extend(["", "## Sources"])
            lines.extend(f"- [{source.label}]({source.reference})" for source in spec.sources)
        return RenderedArtifact(
            content=("\n".join(lines).rstrip() + "\n").encode("utf-8"),
            extension="md",
            media_type="text/markdown; charset=utf-8",
        )


class JsonRenderer(ArtifactRenderer):
    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        payload = spec.model_dump(mode="json")
        payload["resolved_template"] = template.to_public_dict()
        return RenderedArtifact(
            content=(json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
            extension="json",
            media_type="application/json",
        )


class DocxRenderer(ArtifactRenderer):
    """Render the standard_business_brief preset with a memo masthead."""

    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        document = Document(str(template.docx_template)) if template.docx_template else Document()
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        self._configure_styles(document, template)
        self._add_masthead(document, template)
        document.add_heading(spec.title, level=1)
        if spec.subtitle:
            subtitle = document.add_paragraph(spec.subtitle)
            subtitle.style = document.styles["Subtitle"]
        metadata = document.add_paragraph()
        metadata.add_run(f"Template: {template.company_name} | Language: {spec.language}")
        metadata.style = document.styles["Caption"]

        for artifact_section in spec.sections:
            document.add_heading(artifact_section.heading, level=2)
            if artifact_section.body:
                document.add_paragraph(artifact_section.body)
            for bullet in artifact_section.bullets:
                document.add_paragraph(bullet, style="List Bullet")

        for table_spec in spec.tables:
            if table_spec.title:
                document.add_heading(table_spec.title, level=2)
            table = document.add_table(rows=1, cols=len(table_spec.columns))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            table.style = "Table Grid"
            total_width = Inches(6.5)
            column_width = total_width / len(table_spec.columns)
            for index, column in enumerate(table_spec.columns):
                cell = table.rows[0].cells[index]
                cell.width = column_width
                cell.text = column
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _shade_cell(cell, "F2F4F7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row_values in table_spec.rows:
                cells = table.add_row().cells
                for index, value in enumerate(row_values):
                    cells[index].width = column_width
                    cells[index].text = "" if value is None else str(value)
                    cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            document.add_paragraph()

        if spec.sources:
            document.add_heading("Sources", level=2)
            for source in spec.sources:
                document.add_paragraph(f"{source.label}: {source.reference}", style="List Bullet")

        self._add_footer(document)
        with tempfile.TemporaryDirectory(prefix="enterprise-docx-") as directory:
            output_path = Path(directory) / "artifact.docx"
            document.save(output_path)
            Document(output_path)
            content = output_path.read_bytes()
        return RenderedArtifact(
            content=content,
            extension="docx",
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    @staticmethod
    def _configure_styles(document: Document, template: CompanyTemplate) -> None:
        primary = RGBColor.from_string(template.primary_color.removeprefix("#"))
        normal = document.styles["Normal"]
        normal.font.name = template.font_family
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.1
        for name, size, before, after, color in (
            ("Title", 22, 0, 12, primary),
            ("Heading 1", 16, 16, 8, primary),
            ("Heading 2", 13, 12, 6, primary),
            ("Heading 3", 12, 8, 4, primary),
        ):
            style = document.styles[name]
            style.font.name = template.font_family
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
        bullet = document.styles["List Bullet"]
        bullet.font.name = template.font_family
        bullet.font.size = Pt(11)
        bullet.paragraph_format.left_indent = Inches(0.5)
        bullet.paragraph_format.first_line_indent = Inches(-0.25)
        bullet.paragraph_format.space_after = Pt(8)
        bullet.paragraph_format.line_spacing = 1.167

    @staticmethod
    def _add_masthead(document: Document, template: CompanyTemplate) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(template.company_name.upper())
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(template.primary_color.removeprefix("#"))
        paragraph.paragraph_format.space_after = Pt(14)
        properties = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "6")
        bottom.set(qn("w:color"), template.primary_color.removeprefix("#"))
        borders.append(bottom)
        properties.append(borders)

    @staticmethod
    def _add_footer(document: Document) -> None:
        for section in document.sections:
            if section.start_type == WD_SECTION.NEW_PAGE:
                section.footer.is_linked_to_previous = True
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.add_run("Page ")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            paragraph._p.append(field)


class PdfRenderer(ArtifactRenderer):
    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        font_name = _register_pdf_font()
        with tempfile.TemporaryDirectory(prefix="enterprise-pdf-") as directory:
            output_path = Path(directory) / "artifact.pdf"
            document = BaseDocTemplate(
                str(output_path),
                pagesize=letter,
                leftMargin=inch,
                rightMargin=inch,
                topMargin=0.9 * inch,
                bottomMargin=0.8 * inch,
                title=spec.title,
                author=template.company_name,
            )
            frame = Frame(document.leftMargin, document.bottomMargin, document.width, document.height)
            document.addPageTemplates(
                PageTemplate(
                    id="brief",
                    frames=[frame],
                    onPage=lambda canvas, doc: _pdf_header_footer(
                        canvas, doc, template, font_name
                    ),
                )
            )
            styles = _pdf_styles(font_name, template)
            story: list[Any] = [Paragraph(escape(spec.title), styles["ArtifactTitle"])]
            if spec.subtitle:
                story.extend([Paragraph(escape(spec.subtitle), styles["ArtifactSubtitle"]), Spacer(1, 8)])
            for section in spec.sections:
                block: list[Any] = [Paragraph(escape(section.heading), styles["ArtifactH2"])]
                if section.body:
                    block.append(Paragraph(_paragraph_html(section.body), styles["ArtifactBody"]))
                if section.bullets:
                    block.append(
                        ListFlowable(
                            [
                                ListItem(Paragraph(_paragraph_html(item), styles["ArtifactBody"]))
                                for item in section.bullets
                            ],
                            bulletType="bullet",
                            leftIndent=22,
                        )
                    )
                story.append(KeepTogether(block))
            for table_spec in spec.tables:
                if table_spec.title:
                    story.append(Paragraph(escape(table_spec.title), styles["ArtifactH2"]))
                data = [[Paragraph(escape(column), styles["ArtifactTableHeader"]) for column in table_spec.columns]]
                for row in table_spec.rows:
                    data.append(
                        [
                            Paragraph(_paragraph_html("" if value is None else str(value)), styles["ArtifactTableBody"])
                            for value in row
                        ]
                    )
                table = Table(data, repeatRows=1, hAlign="LEFT", colWidths=[document.width / len(table_spec.columns)] * len(table_spec.columns))
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(template.primary_color)),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D0D5DD")),
                            ("LEFTPADDING", (0, 0), (-1, -1), 6),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ("TOPPADDING", (0, 0), (-1, -1), 5),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ]
                    )
                )
                story.extend([table, Spacer(1, 8)])
            if spec.sources:
                story.append(Paragraph("Sources", styles["ArtifactH2"]))
                story.append(
                    ListFlowable(
                        [
                            ListItem(
                                Paragraph(
                                    f"{escape(source.label)}: {escape(source.reference)}",
                                    styles["ArtifactBody"],
                                )
                            )
                            for source in spec.sources
                        ],
                        bulletType="bullet",
                        leftIndent=22,
                    )
                )
            document.build(story)
            PdfReader(output_path)
            content = output_path.read_bytes()
        return RenderedArtifact(content=content, extension="pdf", media_type="application/pdf")


class ArtifactToolRenderer(ArtifactRenderer):
    """Render XLSX/PPTX through the required JavaScript artifact-tool runtime."""

    def __init__(self, node_executable: str, entrypoint: Path | None, timeout_seconds: float) -> None:
        self.node_executable = node_executable
        self.entrypoint = entrypoint
        self.timeout_seconds = timeout_seconds
        self.script_path = Path(__file__).with_name("node_renderer.mjs")

    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        if spec.artifact_type not in {ArtifactFormat.XLSX, ArtifactFormat.PPTX}:
            raise ArtifactRenderingError("artifact-tool renderer received an unsupported format")
        if self.entrypoint is None or not self.entrypoint.is_file():
            raise ArtifactRenderingError(
                "XLSX/PPTX rendering requires ARTIFACT_TOOL_NODE_MODULES with @oai/artifact-tool 2.7.3+."
            )
        with tempfile.TemporaryDirectory(prefix="enterprise-artifact-") as directory:
            build_dir = Path(directory)
            input_path = build_dir / "input.json"
            output_path = build_dir / f"output.{spec.artifact_type.value}"
            preview_dir = build_dir / "preview"
            payload = {
                "spec": spec.model_dump(mode="json"),
                "template": template.to_public_dict(),
            }
            input_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            try:
                completed = subprocess.run(
                    [
                        self.node_executable,
                        str(self.script_path),
                        str(self.entrypoint),
                        str(input_path),
                        str(output_path),
                        str(preview_dir),
                    ],
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    timeout=self.timeout_seconds,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired) as exc:
                raise ArtifactRenderingError("Office artifact renderer could not run") from exc
            if completed.returncode != 0 or not output_path.is_file():
                logger.error(
                    "artifact-tool renderer failed",
                    extra={"format": spec.artifact_type.value, "returncode": completed.returncode},
                )
                raise ArtifactRenderingError("Office artifact renderer failed validation")
            content = output_path.read_bytes()
        if spec.artifact_type is ArtifactFormat.XLSX:
            return RenderedArtifact(
                content=content,
                extension="xlsx",
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        return RenderedArtifact(
            content=content,
            extension="pptx",
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )


class RendererRegistry:
    def __init__(self, renderers: dict[ArtifactFormat, ArtifactRenderer]) -> None:
        missing = set(ArtifactFormat) - set(renderers)
        if missing:
            raise ValueError(f"Missing artifact renderers: {', '.join(sorted(missing))}")
        self._renderers = renderers

    def render(self, spec: ArtifactSpec, template: CompanyTemplate) -> RenderedArtifact:
        return self._renderers[spec.artifact_type].render(spec, template)


def create_renderer_registry(
    *,
    node_executable: str,
    artifact_tool_entrypoint: Path | None,
    timeout_seconds: float,
) -> RendererRegistry:
    artifact_tool = ArtifactToolRenderer(node_executable, artifact_tool_entrypoint, timeout_seconds)
    return RendererRegistry(
        {
            ArtifactFormat.MARKDOWN: MarkdownRenderer(),
            ArtifactFormat.JSON: JsonRenderer(),
            ArtifactFormat.DOCX: DocxRenderer(),
            ArtifactFormat.XLSX: artifact_tool,
            ArtifactFormat.PPTX: artifact_tool,
            ArtifactFormat.PDF: PdfRenderer(),
        }
    )


def _shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _register_pdf_font() -> str:
    candidates = (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            if "ArtifactSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ArtifactSans", str(candidate)))
            return "ArtifactSans"
    return "Helvetica"


def _pdf_styles(font_name: str, template: CompanyTemplate) -> dict[str, ParagraphStyle]:
    styles = getSampleStyleSheet()
    primary = colors.HexColor(template.primary_color)
    return {
        "ArtifactTitle": ParagraphStyle(
            "ArtifactTitle",
            parent=styles["Title"],
            fontName=font_name,
            fontSize=22,
            leading=27,
            textColor=primary,
            alignment=TA_LEFT,
            spaceAfter=12,
        ),
        "ArtifactSubtitle": ParagraphStyle(
            "ArtifactSubtitle",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#475467"),
        ),
        "ArtifactH2": ParagraphStyle(
            "ArtifactH2",
            parent=styles["Heading2"],
            fontName=font_name,
            fontSize=13,
            leading=16,
            textColor=primary,
            spaceBefore=12,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "ArtifactBody": ParagraphStyle(
            "ArtifactBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=10.5,
            leading=14,
            textColor=colors.HexColor("#101828"),
            spaceAfter=6,
        ),
        "ArtifactTableHeader": ParagraphStyle(
            "ArtifactTableHeader",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=8.5,
            leading=10,
            textColor=primary,
        ),
        "ArtifactTableBody": ParagraphStyle(
            "ArtifactTableBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=8.5,
            leading=10,
            textColor=colors.HexColor("#101828"),
        ),
    }


def _paragraph_html(value: str) -> str:
    return escape(value).replace("\n", "<br/>")


def _pdf_header_footer(canvas: Any, document: Any, template: CompanyTemplate, font: str) -> None:
    canvas.saveState()
    canvas.setFont(font, 8)
    canvas.setFillColor(colors.HexColor(template.primary_color))
    canvas.drawString(inch, letter[1] - 0.55 * inch, template.company_name.upper())
    canvas.setStrokeColor(colors.HexColor(template.primary_color))
    canvas.line(inch, letter[1] - 0.62 * inch, letter[0] - inch, letter[1] - 0.62 * inch)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawRightString(letter[0] - inch, 0.45 * inch, f"Page {document.page}")
    canvas.restoreState()
