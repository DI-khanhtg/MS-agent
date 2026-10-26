"""Bounded, offline text extraction for common documents; no OCR or visual analysis."""

import io
from dataclasses import dataclass
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

SUPPORTED = {".txt", ".md", ".csv", ".json", ".docx", ".pptx", ".pdf", ".xlsx"}
_OFFICE = {".docx", ".pptx", ".xlsx"}
_VISUAL_WARNING = (
    "Text only: images, charts and other visual content, if present, are not extracted."
)
_MAX_SOURCE_BYTES = 5_000_000
_MAX_EXPANDED_BYTES = 50_000_000
_MAX_DEPTH = 32


class DocumentExtractionError(ValueError):
    """A document cannot be safely extracted within the supported limits."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    format: str
    unit_count: int | None
    warnings: tuple[str, ...]

    @property
    def metadata(self) -> dict:
        return {
            "format": self.format,
            "method": "text_extraction",
            "unit_count": self.unit_count,
            "coverage": "text_only",
            "warnings": list(self.warnings),
        }


class _Text:
    """Account for every separator before retaining another text fragment."""

    def __init__(self, limit: int):
        self.limit = limit
        self.size = 0
        self.parts: list[str] = []

    def write(self, text: str) -> None:
        if self.size + len(text) > self.limit:
            raise DocumentExtractionError(
                "Extracted document exceeds the text limit; split the file"
            )
        if text:
            self.parts.append(text)
            self.size += len(text)

    def line(self, text: str) -> None:
        if text:
            if self.parts:
                self.write("\n")
            self.write(text)

    def result(self) -> str:
        return "".join(self.parts)


def _preflight_office(content: bytes) -> None:
    with ZipFile(io.BytesIO(content)) as package:
        expanded = 0
        for item in package.infolist():
            if item.flag_bits & 1:
                raise DocumentExtractionError("Encrypted Office documents are unsupported")
            expanded += item.file_size
            if expanded > _MAX_EXPANDED_BYTES:
                raise DocumentExtractionError("Expanded Office document exceeds the size limit")


def _docx_blocks(container, text: _Text, depth: int = 0) -> None:
    if depth > _MAX_DEPTH:
        raise DocumentExtractionError("Document nesting exceeds the supported depth; simplify it")
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            text.line(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                # Merged cells share the same XML; emit each distinct cell once per row.
                seen = set()
                cells = _Text(text.limit - text.size)
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    cell_text = _Text(text.limit - text.size)
                    _docx_blocks(cell, cell_text, depth + 1)
                    if cells.parts:
                        cells.write(" | ")
                    cells.write(cell_text.result())
                text.line(cells.result())


def _extract_docx(content: bytes, text: _Text) -> None:
    document = Document(io.BytesIO(content))
    seen = set()
    # Headers and footers live in separate parts; inherited parts are emitted once.
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header._has_definition and header.part.partname not in seen:
                seen.add(header.part.partname)
                _docx_blocks(header, text)
    _docx_blocks(document, text)
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            if footer._has_definition and footer.part.partname not in seen:
                seen.add(footer.part.partname)
                _docx_blocks(footer, text)


def _pptx_shapes(shapes, text: _Text, depth: int = 0) -> None:
    if depth > _MAX_DEPTH:
        raise DocumentExtractionError("Slide nesting exceeds the supported depth; simplify it")
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _pptx_shapes(shape.shapes, text, depth + 1)
        elif shape.has_table:
            for row in shape.table.rows:
                row_text = _Text(text.limit - text.size)
                has_text = False
                for index, cell in enumerate(row.cells):
                    if index:
                        row_text.write(" | ")
                    if not cell.is_spanned:
                        row_text.write(cell.text)
                        has_text |= bool(cell.text.strip())
                if has_text:
                    text.line(row_text.result())
        elif shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.text.strip():
                    text.line(paragraph.text)


def _extract_pptx(content: bytes, text: _Text, warnings: list[str]) -> int:
    deck = Presentation(io.BytesIO(content))
    if len(deck.slides) > 200:
        raise DocumentExtractionError(
            "Presentations over 200 slides are unsupported; split the file"
        )
    empty = 0
    for index, slide in enumerate(deck.slides, 1):
        slide_text = _Text(text.limit - text.size)
        _pptx_shapes(slide.shapes, slide_text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                slide_text.line("Speaker notes:")
                for paragraph in notes.paragraphs:
                    if paragraph.text.strip():
                        slide_text.line(paragraph.text)
        if slide_text.parts:
            text.line(f"Slide {index}")
            text.line(slide_text.result())
        else:
            empty += 1
    if empty:
        warnings.append(
            f"{empty} of {len(deck.slides)} PowerPoint slides have no readable text; "
            "images and charts are not read."
        )
    return len(deck.slides)


def _extract_pdf(content: bytes, text: _Text, warnings: list[str]) -> int:
    document = PdfReader(io.BytesIO(content))
    if document.is_encrypted:
        raise DocumentExtractionError("Encrypted PDFs are unsupported; provide an unencrypted copy")
    count = len(document.pages)
    if count > 200:
        raise DocumentExtractionError("PDFs over 200 pages are unsupported; split the file")
    empty = 0
    for index, page in enumerate(document.pages, 1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            text.line(f"Page {index}")
            text.line(page_text)
        else:
            empty += 1
    if empty:
        warnings.append(
            f"{empty} of {count} PDF pages have no selectable text; "
            "scanned or image-only pages require OCR, which is not performed."
        )
    return count


def _extract_xlsx(content: bytes, text: _Text, warnings: list[str]) -> int:
    warnings.append(
        "Only cached formula values are read; unavailable cached values are omitted. "
        "Formulas are not evaluated."
    )
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        empty = 0
        for sheet in workbook:
            sheet_text = _Text(text.limit - text.size)
            has_values = False
            for row in sheet.iter_rows(values_only=True):
                row_text = _Text(sheet_text.limit - sheet_text.size)
                for index, value in enumerate(row):
                    if index:
                        row_text.write(" | ")
                    if value is not None:
                        value_text = str(value)
                        row_text.write(value_text)
                        has_values |= bool(value_text.strip())
                # Account for empty rows too, bounding worksheets with inflated dimensions.
                sheet_text.write("\n")
                sheet_text.write(row_text.result())
            if has_values:
                text.line(f"Sheet: {sheet.title}")
                text.write(sheet_text.result())
            else:
                empty += 1
        if empty:
            warnings.append(
                f"{empty} of {len(workbook.worksheets)} worksheets have no readable cached "
                "values; they may be empty or contain formulas without cached results."
            )
        return len(workbook.worksheets)
    finally:
        workbook.close()


def extract_document(
    content: bytes, suffix: str, max_chars: int = 100_000
) -> ExtractedDocument:
    """Extract selectable text without executing macros or fetching external resources.

    Source files are capped at 5 MB, expanded Office ZIPs at 50 MB, PDF/PPTX at
    200 pages/slides, and text is rejected as soon as the requested limit is exceeded.
    Legacy binary Office files must be converted before reading.
    """
    if not isinstance(max_chars, int) or isinstance(max_chars, bool) or max_chars <= 0:
        raise DocumentExtractionError("max_chars must be a positive integer")
    suffix = suffix.lower()
    if suffix in {".doc", ".ppt", ".xls"}:
        raise DocumentExtractionError("Convert legacy Office files to DOCX, PPTX or XLSX first")
    if suffix not in SUPPORTED:
        raise DocumentExtractionError("Unsupported document format")
    if len(content) > _MAX_SOURCE_BYTES:
        raise DocumentExtractionError("Source document exceeds the 5 MB size limit; split the file")
    text = _Text(max_chars)
    warnings: list[str] = []
    count = None
    try:
        if suffix in _OFFICE:
            _preflight_office(content)
        if suffix == ".docx":
            _extract_docx(content, text)
        elif suffix == ".pptx":
            count = _extract_pptx(content, text, warnings)
        elif suffix == ".pdf":
            count = _extract_pdf(content, text, warnings)
        elif suffix == ".xlsx":
            count = _extract_xlsx(content, text, warnings)
        else:
            decoded = content.decode("utf-8-sig")
            if any(ord(char) < 32 and char not in "\t\n\r" for char in decoded):
                raise DocumentExtractionError(
                    "Text documents must contain UTF-8 text, not binary data"
                )
            text.write(decoded)
    except DocumentExtractionError:
        raise
    except Exception:
        raise DocumentExtractionError(
            "Document could not be extracted; check that it is a valid, unencrypted "
            "file matching its extension, or re-export it in a supported format"
        ) from None
    if suffix in _OFFICE or suffix == ".pdf":
        warnings.append(_VISUAL_WARNING)
    return ExtractedDocument(text.result(), suffix[1:], count, tuple(warnings))


def extract_text(content: bytes, suffix: str, max_chars: int = 100_000) -> str:
    """Compatibility wrapper for callers which only need the extracted text."""
    return extract_document(content, suffix, max_chars).text
