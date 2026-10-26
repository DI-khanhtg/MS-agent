"""Small real document fixtures shared by adapter, API and browser checks."""

import io

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas


@pytest.fixture(scope="session")
def document_files():
    files = {}
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("Word project requirements")
    document.save(stream)
    files["docx"] = ("requirements.docx", stream.getvalue(), "Word project requirements")
    stream = io.BytesIO()
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(6), Inches(1)
    ).text = "PowerPoint project objectives"
    slides.save(stream)
    files["pptx"] = ("objectives.pptx", stream.getvalue(), "PowerPoint project objectives")
    stream = io.BytesIO()
    pdf = Canvas(stream)
    pdf.drawString(40, 750, "PDF project schedule")
    pdf.save()
    files["pdf"] = ("schedule.pdf", stream.getvalue(), "PDF project schedule")
    stream = io.BytesIO()
    book = Workbook()
    book.active.append(["Excel project budget", 100])
    book.save(stream)
    book.close()
    files["xlsx"] = ("budget.xlsx", stream.getvalue(), "Excel project budget")
    return files
