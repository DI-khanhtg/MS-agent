import asyncio
import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from enterprise_agent.agent.runtime import MicrosoftAgentRuntime
from enterprise_agent.api.app import create_app
from enterprise_agent.artifacts.models import ArtifactSpec
from enterprise_agent.artifacts.tools import ArtifactTools
from enterprise_agent.config import Settings
from enterprise_agent.security.identity import Principal, RequestContext


class ExportAgent:
    def __init__(self, tools, formats, *, text="Created.", fail=False):
        self.tools, self.formats, self.text, self.fail = tools, formats, text, fail
        self.calls = 0

    def create_session(self, **kwargs):
        return {}

    def run(self, message, *, session, stream=False):
        self.calls += 1
        return self.updates() if stream else self.reply()

    async def reply(self):
        record = None
        for kind in self.formats:
            # Framework tools may run in child tasks. Receipts must still reach the caller.
            record = await asyncio.create_task(
                self.tools.create_artifact(
                    ArtifactSpec(
                        artifact_type=kind,
                        title="Chat export",
                        sections=[{"heading": "Analysis", "body": "Verified export content"}],
                    )
                )
            )
        if self.fail:
            raise ConnectionError("private-provider-diagnostic")
        text = self.text.replace("{download_url}", record["download_url"]) if record else self.text
        return SimpleNamespace(text=text, usage_details={}, finish_reason="stop")

    async def updates(self):
        yield await self.reply()


def export_app(tmp_path, formats, **kwargs):
    settings = Settings(
        _env_file=None,
        runtime_mode="demo",
        gemini_max_retries=2,
        artifact_storage_path=tmp_path / "artifacts",
        audit_log_path=tmp_path / "audit.jsonl",
        artifact_tool_node_modules=tmp_path / "no-private-node-runtime",
    )
    app = create_app(settings=settings)
    agent = ExportAgent(ArtifactTools(app.state.artifact_service), formats, **kwargs)
    app.state.agent_runtime = MicrosoftAgentRuntime(settings, agent=agent)
    return TestClient(app), agent


def validate_export(kind, content):
    stream = io.BytesIO(content)
    if kind == "docx":
        assert "Verified export content" in "\n".join(p.text for p in Document(stream).paragraphs)
    elif kind == "pptx":
        assert "Verified export content" in "\n".join(
            shape.text
            for slide in Presentation(stream).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    elif kind == "xlsx":
        book = load_workbook(stream, read_only=True)
        try:
            assert any("Verified export content" in row for row in book.active.values)
        finally:
            book.close()
    elif kind == "pdf":
        assert "Verified export content" in "\n".join(
            p.extract_text() for p in PdfReader(stream).pages
        )
    elif kind == "json":
        assert json.loads(content)["sections"][0]["body"] == "Verified export content"
    else:
        assert "Verified export content" in content.decode("utf-8")


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx", "pdf", "markdown", "json"])
@pytest.mark.parametrize("stream", [False, True])
@pytest.mark.parametrize("text", ["", "Created."])
def test_chat_returns_a_real_download_even_when_model_omits_the_link(tmp_path, kind, stream, text):
    client, agent = export_app(tmp_path, [kind], text=text)
    response = client.post(
        "/api/chat/stream" if stream else "/api/chat",
        json={"message": "Create a test document", "session_id": "export-session"},
    )
    assert response.status_code == 200
    if stream:
        assert "event: done" in response.text
        assert "event: error" not in response.text
        answer = "".join(
            json.loads(block.split("data: ", 1)[1])["text"]
            for block in response.text.split("\n\n")
            if block.startswith("event: delta\n")
        )
    else:
        answer = response.json()["message"]
    records = client.get("/api/artifacts").json()
    assert len(records) == 1 and records[0]["artifact_type"] == kind
    link = records[0]["download_url"]
    assert answer.count(link) == 1
    download = client.get(link)
    assert download.status_code == 200
    validate_export(kind, download.content)
    saved = client.get("/api/sessions/export-session").json()
    assert link in saved["messages"][-1]["text"]
    assert agent.calls == 1


def test_files_created_before_provider_failure_remain_downloadable_in_chat(tmp_path):
    client, agent = export_app(tmp_path, ["docx"], fail=True)
    response = client.post("/api/chat/stream", json={"message": "Create a Word report"})
    records = client.get("/api/artifacts").json()
    assert len(records) == 1
    assert records[0]["download_url"] in response.text
    assert "event: error" in response.text
    assert "event: done" not in response.text
    assert "private-provider-diagnostic" not in response.text
    assert agent.calls == 1


def test_multiple_formats_get_separate_links_and_do_not_leak_into_the_next_turn(tmp_path):
    client, agent = export_app(tmp_path, ["docx", "pptx"], text="Here are the files.")
    response = client.post("/api/chat", json={"message": "Create Word and PowerPoint"})
    assert response.status_code == 200
    records = client.get("/api/artifacts").json()
    assert len(records) == 2
    for record in records:
        assert response.json()["message"].count(record["download_url"]) == 1
    agent.formats = []
    second = client.post("/api/chat", json={"message": "Hello"})
    assert "/api/artifacts/" not in second.json()["message"]


def test_history_preserves_verified_file_links_while_redacting_other_pii(tmp_path, monkeypatch):
    from uuid import UUID

    artifact_id = "12345678-1234-4234-8234-123456789012"
    monkeypatch.setattr("enterprise_agent.artifacts.store.uuid4", lambda: UUID(artifact_id))
    client, agent = export_app(tmp_path, ["docx"], text="Contact: 0901234567")
    response = client.post("/api/chat", json={"message": "Export", "session_id": "pii-links"})
    assert response.status_code == 200
    link = f"/api/artifacts/{artifact_id}/download"
    saved = client.get("/api/sessions/pii-links").json()["messages"][-1]["text"]
    assert link in saved
    assert "0901234567" not in saved
    assert "[REDACTED]" in saved
    agent.formats = []
    agent.text = f"[Existing file]({link})"
    client.post("/api/chat", json={"message": "Show that link again", "session_id": "pii-links"})
    assert link in client.get("/api/sessions/pii-links").json()["messages"][-1]["text"]


@pytest.mark.parametrize("stream", [False, True])
def test_verified_link_already_in_the_answer_is_not_duplicated(tmp_path, stream):
    client, _ = export_app(tmp_path, ["docx"], text="[Download]({download_url})")
    response = client.post(
        "/api/chat/stream" if stream else "/api/chat", json={"message": "Export"}
    )
    link = client.get("/api/artifacts").json()[0]["download_url"]
    text = response.text if stream else response.json()["message"]
    assert text.count(link) == 1


@pytest.mark.parametrize("text", ["File: {download_url}", "`[Download]({download_url})`"])
def test_literal_urls_and_code_are_followed_by_a_clickable_download_link(tmp_path, text):
    client, _ = export_app(tmp_path, ["docx"], text=text)
    response = client.post("/api/chat", json={"message": "Export"})
    link = client.get("/api/artifacts").json()[0]["download_url"]
    assert f"[Tải DOCX]({link})" in response.json()["message"]


def test_pii_filter_does_not_whitelist_an_unverified_download_url(tmp_path):
    link = "/api/artifacts/12345678-1234-4234-8234-123456789012/download"
    client, _ = export_app(tmp_path, [], text=f"[Unknown file]({link})")
    client.post("/api/chat", json={"message": "Show file", "session_id": "unverified"})
    saved = client.get("/api/sessions/unverified").json()["messages"][-1]["text"]
    assert link not in saved
    assert "[REDACTED]" in saved


@pytest.mark.asyncio
async def test_concurrent_exports_keep_receipts_isolated_by_user(tmp_path):
    client, _ = export_app(tmp_path, ["docx"], text="Created.")
    runtime = client.app.state.agent_runtime
    owners = [
        Principal("tenant", name, name, None, frozenset(), "frontend", "test-token")
        for name in ("alice", "bob")
    ]
    results = await asyncio.gather(
        *[runtime.chat("Export", "same-session", RequestContext(owner)) for owner in owners]
    )
    store = client.app.state.artifact_service.store
    links = [store.list(owner)[0].download_url for owner in owners]
    assert links[0] != links[1]
    for index, result in enumerate(results):
        assert links[index] in result.text
        assert links[1 - index] not in result.text
