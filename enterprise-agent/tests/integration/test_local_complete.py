import io
import json
import wave
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from enterprise_agent.agent.runtime import AgentResult
from enterprise_agent.api.app import create_app
from enterprise_agent.config import Settings
from enterprise_agent.security.identity import Principal
from enterprise_agent.voice.transcriber import Transcript, validate_wav


def local_settings(tmp_path: Path, **overrides) -> Settings:
    root = tmp_path / "documents" / "Project-X"
    root.mkdir(parents=True, exist_ok=True)
    (root / "requirements.md").write_text(
        "# Requirements\nUse Entra ID. Support 50 users.", encoding="utf-8"
    )
    (root / "data.csv").write_text("metric,value\nusers,50\n", encoding="utf-8")
    values = dict(
        runtime_mode="demo",
        storage_backend="database",
        enabled_tool_groups="local_read,artifact_generation",
        database_path=tmp_path / "agent.db",
        local_documents_root=root.parent,
        artifact_storage_path=tmp_path / "artifacts",
        audit_log_path=tmp_path / "audit.jsonl",
    )
    values.update(overrides)
    return Settings(_env_file=None, **values)


@pytest.mark.parametrize("fmt", ["docx", "pdf", "xlsx", "pptx", "markdown", "json"])
def test_offline_folder_to_downloadable_artifact(tmp_path, monkeypatch, fmt):
    async def no_network(*args, **kwargs):
        pytest.fail("Offline demo attempted a network request")

    monkeypatch.setattr(httpx.AsyncClient, "request", no_network)
    settings = local_settings(tmp_path)
    with TestClient(create_app(settings=settings)) as client:
        response = client.post(
            "/api/chat",
            json={
                "message": f"Read folder Project-X and create proposal {fmt}",
                "session_id": "demo",
            },
        )
        assert response.status_code == 200, response.text
        assert "demo" in response.json()["message"]
        tools = response.json()["tool_calls"]
        assert [t["name"] for t in tools] == [
            "list_local_documents",
            "read_local_document",
            "read_local_document",
            "create_artifact",
        ]
        item = client.get("/api/artifacts").json()[0]
        artifact = client.get(item["download_url"])
        assert artifact.status_code == 200
        content = io.BytesIO(artifact.content)
        if fmt == "docx":
            assert "Proposal" in Document(content).paragraphs[1].text
        elif fmt == "pdf":
            assert len(PdfReader(content).pages) > 0
        elif fmt == "xlsx":
            assert "Sources" in load_workbook(content).sheetnames
        elif fmt == "pptx":
            assert len(Presentation(content).slides) >= 5
        elif fmt == "json":
            assert json.loads(artifact.content)["metadata"]["runtime"] == "demo"
        else:
            assert "local://Project-X/requirements.md" in artifact.text
    with TestClient(create_app(settings=settings)) as restarted:
        history = restarted.get("/api/sessions/demo").json()
        assert history["turns"] == 1
        assert history["messages"][-1]["tool_calls"][-1]["name"] == "create_artifact"
        assert restarted.get("/api/artifacts").json()[0]["artifact_id"] == item["artifact_id"]


def test_stream_has_tools_and_persists_history(tmp_path):
    with TestClient(create_app(settings=local_settings(tmp_path))) as client:
        response = client.post(
            "/api/chat/stream",
            json={
                "message": "Read folder Project-X and create proposal docx",
                "session_id": "stream",
            },
        )
        assert "event: done" in response.text
        assert "create_artifact" in response.text
        history = client.get("/api/sessions/stream").json()
        assert len(history["messages"][-1]["tool_calls"]) == 4


def wav_bytes(rate=16000):
    data = io.BytesIO()
    with wave.open(data, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(b"\0\0" * rate)
    return data.getvalue()


def test_voice_adapter_and_disabled_provider(tmp_path):
    class FakeSpeech:
        async def transcribe(self, audio, language):
            _, duration = validate_wav(audio)
            return Transcript("Create a proposal", language, duration, provider="test")

    settings = local_settings(tmp_path)
    with TestClient(create_app(settings=settings, transcriber=FakeSpeech())) as client:
        response = client.post(
            "/api/voice/transcribe", files={"audio": ("command.wav", wav_bytes(), "audio/wav")}
        )
        assert response.json()["text"] == "Create a proposal"
        assert response.json()["provider"] == "test"
    with TestClient(create_app(settings=settings)) as client:
        assert (
            client.post(
                "/api/voice/transcribe", files={"audio": ("command.wav", wav_bytes(), "audio/wav")}
            ).status_code
            == 503
        )
        assert (
            client.post(
                "/api/voice/transcribe", files={"audio": ("bad.wav", wav_bytes(8000), "audio/wav")}
            ).status_code
            == 422
        )


def test_rate_limit_and_cross_origin_protection(tmp_path):
    with TestClient(
        create_app(settings=local_settings(tmp_path, rate_limit_per_minute=2))
    ) as client:
        assert client.get("/api/config").status_code == 200
        assert client.get("/api/config").status_code == 200
        assert client.get("/api/config").status_code == 429
        response = client.post(
            "/api/chat", headers={"Origin": "https://evil.example"}, json={"message": "hello"}
        )
        assert response.status_code == 403


def test_local_failures_and_metadata_only_audit(tmp_path):
    settings = local_settings(tmp_path)
    with TestClient(create_app(settings=settings)) as client:
        assert (
            client.post(
                "/api/chat", json={"message": "Read folder ../outside and create proposal"}
            ).status_code
            == 422
        )
        assert (
            client.post(
                "/api/chat", json={"message": "Read folder missing and create proposal"}
            ).status_code
            == 422
        )
        client.post("/api/chat", json={"message": "private unique user content"})
        assert "private unique user content" not in settings.audit_log_path.read_text()
        assert client.get("/").status_code == 200
        assert client.get("/static/app.js").status_code == 200


def test_chunked_body_limit_cannot_be_bypassed(tmp_path):
    with TestClient(create_app(settings=local_settings(tmp_path))) as client:
        response = client.post(
            "/api/voice/transcribe",
            content=iter([b"x" * 1_500_000] * 2),
            headers={"Content-Type": "application/octet-stream"},
        )
        assert response.status_code == 413


def test_durable_history_is_isolated_by_verified_user_and_delete(tmp_path):
    class Validator:
        async def validate(self, token):
            return Principal(
                "tenant", token, token, None, frozenset({"access_as_user"}), "client", token
            )

    class Runtime:
        async def chat(self, message, session_id, context):
            return AgentResult("answer", [], {"input_token_count": 1000, "output_token_count": 500})

    settings = local_settings(
        tmp_path,
        auth_enabled=True,
        enabled_tool_groups="artifact_generation",
        entra_api_audience="client",
        azure_tenant_id="tenant",
        input_price_per_million=1,
        output_price_per_million=2,
    )
    app = create_app(settings=settings, token_validator=Validator(), runtime=Runtime())
    with TestClient(app) as client:
        user_a = {"Authorization": "Bearer user-a"}
        user_b = {"Authorization": "Bearer user-b"}
        assert (
            client.post(
                "/api/chat", json={"message": "hello", "session_id": "same"}, headers=user_a
            ).status_code
            == 200
        )
        assert client.get("/api/sessions/same", headers=user_b).status_code == 404
        assert client.get("/api/sessions", headers=user_b).json() == []
        assert client.delete("/api/sessions/same", headers=user_b).status_code == 404
        assert (
            client.post(
                "/api/chat", json={"message": "user b", "session_id": "same"}, headers=user_b
            ).status_code
            == 200
        )
        assert (
            client.get("/api/sessions/same", headers=user_a).json()["messages"][0]["text"]
            == "hello"
        )
        assert (
            client.get("/api/sessions/same", headers=user_b).json()["messages"][0]["text"]
            == "user b"
        )
        assert client.get("/api/metrics", headers=user_a).status_code == 403
    with TestClient(
        create_app(settings=settings, token_validator=Validator(), runtime=Runtime())
    ) as client:
        assert client.get("/api/sessions/same", headers=user_a).json()["turns"] == 1
        assert client.delete("/api/sessions/same", headers=user_a).status_code == 204
        assert client.get("/api/sessions/same", headers=user_a).status_code == 404
        assert client.get("/api/sessions/same", headers=user_b).status_code == 200


def test_usage_cost_is_explicit_and_audit_records_egress(tmp_path):
    settings = local_settings(tmp_path, input_price_per_million=1, output_price_per_million=2)
    with TestClient(create_app(settings=settings)) as client:
        client.post("/api/chat", json={"message": "Read folder Project-X and create proposal"})
        events = [json.loads(line) for line in settings.audit_log_path.read_text().splitlines()]
        assert any(item["event"] == "egress_decision" for item in events)
        assert any(item["event"] == "artifact_created" for item in events)
        metrics = client.get("/api/metrics").json()
        assert metrics["estimated_model_cost"] == 0
        assert metrics["requests"] == 1
        assert metrics["tool_calls"] == 4


def test_long_source_pptx_is_paginated_and_reports_partial_coverage(tmp_path):
    settings = local_settings(tmp_path)
    (settings.local_documents_root / "Project-X" / "long.md").write_text("Details " * 600)
    with TestClient(create_app(settings=settings)) as client:
        response = client.post(
            "/api/chat", json={"message": "Read folder Project-X, proposal pptx"}
        )
        assert response.status_code == 200, response.text
        assert "Coverage" in response.json()["message"]
        artifact = client.get("/api/artifacts").json()[0]
        assert client.get(artifact["download_url"]).status_code == 200


def test_production_rejects_demo_and_unauthenticated_start(tmp_path):
    with pytest.raises(ValueError, match="Staging/production"):
        create_app(settings=local_settings(tmp_path, app_env="production"))
