"""Real document fixtures exercise extraction, coverage and resource limits."""

import io
import struct
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas


def reader():
    from enterprise_agent.tools import document_reader

    return document_reader


def saved(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_with_text_and_blank_page():
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(72, 720, "Selectable report")
    canvas.showPage()
    canvas.showPage()
    canvas.save()
    return output.getvalue()


@pytest.mark.parametrize("suffix", [".txt", ".MD", ".csv", ".json"])
def test_utf8_text_is_preserved_and_suffix_normalized(suffix):
    result = reader().extract_document(b"\xef\xbb\xbfHello\nworld", suffix)
    assert result.text == "Hello\nworld"
    assert result.metadata == {
        "format": suffix[1:].lower(),
        "method": "text_extraction",
        "unit_count": None,
        "coverage": "text_only",
        "warnings": [],
    }
    assert reader().extract_text(b"hello", suffix) == "hello"


def test_docx_preserves_body_table_order_nested_cells_headers_and_footers():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Report header"
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell"
    table.cell(0, 1).text = "Right cell"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    document.add_paragraph("After table")
    document.sections[0].footer.paragraphs[0].text = "Report footer"
    result = reader().extract_document(saved(document), ".docx")
    positions = [result.text.index(t) for t in (
        "Report header", "Before table", "Left cell", "Nested cell",
        "Right cell", "After table", "Report footer",
    )]
    assert positions == sorted(positions)
    assert result.format == "docx"
    assert result.unit_count is None
    assert any("visual" in warning.lower() for warning in result.warnings)


def test_pptx_extracts_grouped_text_tables_notes_and_slide_order():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Grouped heading"
    table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Revenue"
    slide.notes_slide.notes_text_frame.text = "Speaker detail"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Second slide"
    result = reader().extract_document(saved(deck), ".PPTX")
    positions = [result.text.index(t) for t in (
        "Slide 1", "Grouped heading", "Product", "Revenue", "Speaker detail",
        "Slide 2", "Second slide",
    )]
    assert positions == sorted(positions)
    assert result.unit_count == 2
    assert "1" not in result.text.splitlines()  # Notes slide-number placeholder is excluded.
    assert any("visual" in warning.lower() for warning in result.warnings)


def test_pdf_labels_text_pages_and_reports_empty_page_coverage():
    result = reader().extract_document(pdf_with_text_and_blank_page(), ".PDF")
    assert "Page 1" in result.text
    assert "Selectable report" in result.text
    assert "Page 2" not in result.text
    assert result.metadata["unit_count"] == 2
    assert any("OCR" in warning for warning in result.warnings)


def test_images_only_or_blank_pdf_stays_empty_and_requests_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    result = reader().extract_document(output.getvalue(), ".pdf")
    assert result.text == ""
    assert result.unit_count == 1
    assert any("OCR" in warning for warning in result.warnings)


def test_image_only_pdf_preserves_empty_text_without_ocr():
    image = b"P6\n1 1\n255\n\xff\x00\x00"  # A real one-pixel RGB PPM image.
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawImage(ImageReader(io.BytesIO(image)), 0, 0, width=200, height=200)
    canvas.showPage()
    canvas.save()
    result = reader().extract_document(output.getvalue(), ".pdf")
    assert result.text == ""
    assert result.unit_count == 1
    assert result.metadata["coverage"] == "text_only"
    assert any("OCR" in warning for warning in result.warnings)


def test_chart_only_pptx_is_empty_and_reports_unreadable_slide():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (3, 7))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), data
    )
    result = reader().extract_document(saved(deck), ".pptx")
    assert result.unit_count == 1
    assert result.text == ""
    assert any("1 of 1" in warning and "no readable text" in warning for warning in result.warnings)
    assert any("charts" in warning for warning in result.warnings)


@pytest.mark.parametrize("kind", ["picture", "empty_table", "whitespace"])
def test_pptx_without_actual_text_does_not_emit_slide_labels(kind):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if kind == "picture":
        image = io.BytesIO()
        Image.new("RGB", (1, 1), color="red").save(image, "PNG")
        image.seek(0)
        slide.shapes.add_picture(image, 0, 0)
    elif kind == "empty_table":
        slide.shapes.add_table(1, 2, 0, 0, Inches(4), Inches(1))
    else:
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "   "
    result = reader().extract_document(saved(deck), ".pptx")
    assert result.text == ""
    assert result.unit_count == 1
    assert any("no readable text" in warning for warning in result.warnings)


def test_pptx_preserves_actual_slide_number_for_notes_only_after_empty_slide():
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "Speaker detail"
    result = reader().extract_document(saved(deck), ".pptx")
    assert result.text == "Slide 2\nSpeaker notes:\nSpeaker detail"
    assert result.unit_count == 2
    assert any("1 of 2" in warning and "no readable text" in warning for warning in result.warnings)


def test_xlsx_extracts_sheet_order_and_warns_about_unavailable_formula_values():
    workbook = Workbook()
    workbook.active.title = "Revenue"
    workbook.active.append(["Quarter", "Value"])
    workbook.active.append(["Q1", 7])
    workbook.active.append(["Total", "=SUM(B2:B2)"])
    workbook.create_sheet("Notes").append(["Unaudited"])
    result = reader().extract_document(saved(workbook), ".xlsx")
    assert "Sheet: Revenue" in result.text
    assert "Quarter | Value" in result.text
    assert "Q1 | 7" in result.text
    assert result.text.index("Total") < result.text.index("Sheet: Notes")
    assert "=SUM" not in result.text
    assert result.unit_count == 2
    assert any("formula" in warning.lower() for warning in result.warnings)


@pytest.mark.parametrize("kind", ["empty", "formula", "whitespace"])
def test_xlsx_without_readable_values_does_not_emit_sheet_labels(kind):
    workbook = Workbook()
    if kind == "formula":
        workbook.active.append(["=1+1", "=2+2"])
    elif kind == "whitespace":
        workbook.active.append(["   ", "\t"])
    result = reader().extract_document(saved(workbook), ".xlsx")
    assert result.text == ""
    assert result.unit_count == 1
    assert any("1 of 1" in warning and "no readable" in warning for warning in result.warnings)


def test_xlsx_reports_empty_sheets_while_retaining_zero_and_false_values():
    workbook = Workbook()
    workbook.active.title = "Empty"
    workbook.create_sheet("Values").append([0, False])
    result = reader().extract_document(saved(workbook), ".xlsx")
    assert result.text == "Sheet: Values\n0 | False"
    assert result.unit_count == 2
    assert any("1 of 2" in warning and "no readable" in warning for warning in result.warnings)


@pytest.mark.parametrize("suffix", [".doc", ".ppt", ".xls", ".exe"])
def test_unsupported_formats_are_rejected_without_decoding_binary(suffix):
    with pytest.raises(reader().DocumentExtractionError, match="[Uu]nsupported|[Cc]onvert"):
        reader().extract_document(b"\xff\x00private", suffix)


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".pdf", ".xlsx", ".txt"])
def test_corrupt_or_binary_content_has_actionable_nonleaking_error(suffix):
    with pytest.raises(reader().DocumentExtractionError) as error:
        reader().extract_document(b"\xff\x00private-content", suffix)
    assert "private-content" not in str(error.value)
    assert str(error.value)


def test_source_byte_limit_rejects_before_decoding():
    with pytest.raises(reader().DocumentExtractionError, match="size|5 MB"):
        reader().extract_document(b"a" * 5_000_001, ".txt")


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".xlsx"])
def test_office_zip_expansion_limit_is_checked_before_parsing(suffix):
    output = io.BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("oversized.xml", b"a" * 50_000_001)
    with pytest.raises(reader().DocumentExtractionError, match="Expanded|expanded"):
        reader().extract_document(output.getvalue(), suffix)


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".xlsx"])
def test_encrypted_office_zip_entries_are_rejected(suffix):
    output = io.BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("content.xml", "encrypted-content-placeholder")
    content = bytearray(output.getvalue())
    # Set the actual ZIP encrypted flags; preflight must reject before reading the payload.
    struct.pack_into("<H", content, 6, 1)
    central = content.index(b"PK\x01\x02")
    struct.pack_into("<H", content, central + 8, 1)
    with pytest.raises(reader().DocumentExtractionError, match="[Ee]ncrypted"):
        reader().extract_document(bytes(content), suffix)


@pytest.mark.parametrize("kind", ["text", "docx", "pptx", "pdf", "xlsx"])
def test_extracted_character_limit_covers_all_formats(kind):
    if kind == "text":
        content, suffix = b"a" * 101, ".txt"
    elif kind == "docx":
        document = Document()
        document.add_paragraph("a" * 101)
        content, suffix = saved(document), ".docx"
    elif kind == "pptx":
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "a" * 101
        content, suffix = saved(document), ".pptx"
    elif kind == "pdf":
        content, suffix = pdf_with_text_and_blank_page(), ".pdf"
    else:
        document = Workbook()
        document.active.append(["a" * 101])
        content, suffix = saved(document), ".xlsx"
    with pytest.raises(reader().DocumentExtractionError, match="text limit"):
        reader().extract_document(content, suffix, max_chars=10)


def test_pdf_encryption_is_rejected_even_when_empty_password_opens_it():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(reader().DocumentExtractionError, match="[Ee]ncrypted"):
        reader().extract_document(output.getvalue(), ".pdf")


@pytest.mark.parametrize("suffix", [".pdf", ".pptx"])
def test_page_and_slide_count_is_bounded(suffix):
    if suffix == ".pdf":
        writer = PdfWriter()
        for _ in range(201):
            writer.add_blank_page(width=100, height=100)
        output = io.BytesIO()
        writer.write(output)
        content = output.getvalue()
    else:
        document = Presentation()
        for _ in range(201):
            document.slides.add_slide(document.slide_layouts[6])
        content = saved(document)
    with pytest.raises(reader().DocumentExtractionError, match="200"):
        reader().extract_document(content, suffix)


@pytest.mark.parametrize("max_chars", [0, -1, True, "10"])
def test_invalid_character_limits_are_rejected(max_chars):
    with pytest.raises(reader().DocumentExtractionError, match="max_chars"):
        reader().extract_document(b"hello", ".txt", max_chars=max_chars)


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_deeply_nested_document_structures_fail_with_actionable_bound(suffix):
    if suffix == ".docx":
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        for _ in range(34):
            cell = cell.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Deep text"
    else:
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        for _ in range(34):
            group = group.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Deep text"
    with pytest.raises(reader().DocumentExtractionError, match="nesting|depth"):
        reader().extract_document(saved(document), suffix)


def test_character_limit_counts_separators_between_document_parts():
    document = Document()
    document.add_paragraph("abc")
    document.add_paragraph("def")
    content = saved(document)
    assert reader().extract_document(content, ".docx", max_chars=7).text == "abc\ndef"
    with pytest.raises(reader().DocumentExtractionError, match="text limit"):
        reader().extract_document(content, ".docx", max_chars=6)
