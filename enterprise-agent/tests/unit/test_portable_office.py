import io

from openpyxl import load_workbook
from pptx import Presentation

from enterprise_agent.artifacts.models import ArtifactSpec
from enterprise_agent.artifacts.office import PptxRenderer, XlsxRenderer
from enterprise_agent.artifacts.templates import CompanyTemplate

TEMPLATE = CompanyTemplate("default", "Enterprise", "#164E63", "#0EA5E9", "Calibri")


def test_dense_tables_and_references_remain_editable_and_paginated():
    spec = ArtifactSpec(
        artifact_type="pptx",
        title="Dense data",
        sections=[{"heading": "Overview", "body": "Verified data"}],
        tables=[
            {
                "columns": ["One", "Two", "Three", "Four", "Five", "Six"],
                "rows": [["a" * 240] * 6 for _ in range(7)],
            }
        ],
        sources=[{"label": "Long source", "reference": "workiq:" + "a" * 1900}],
    )
    result = PptxRenderer().render(spec, TEMPLATE)
    deck = Presentation(io.BytesIO(result.content))
    tables = [shape.table for page in deck.slides for shape in page.shapes if shape.has_table]
    assert sum(len(table.rows) - 1 for table in tables) == 7
    assert len(tables) > 1
    for page in deck.slides:
        for shape in page.shapes:
            if shape.has_table:
                assert shape.top + shape.height < deck.slide_height
    source_text = "".join(
        shape.text
        for page in deck.slides
        for shape in page.shapes
        if shape.has_text_frame and shape.text.startswith(("workiq:", "Long source", "a"))
    )
    assert "workiq:" in source_text
    assert source_text.count("a") >= 1900


def test_long_spreadsheet_sections_are_wrapped_without_losing_content():
    body = "Requirement " * 500
    spec = ArtifactSpec(
        artifact_type="xlsx", title="Report", sections=[{"heading": "Requirements", "body": body}]
    )
    result = XlsxRenderer().render(spec, TEMPLATE)
    sheet = load_workbook(io.BytesIO(result.content)).active
    chunks = [row[0] for row in sheet.iter_rows(min_row=3)]
    assert "".join(cell.value for cell in chunks) == body.strip()
    assert all(cell.alignment.wrap_text for cell in chunks)
